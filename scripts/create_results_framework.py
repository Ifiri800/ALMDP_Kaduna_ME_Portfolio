from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# ============================================================
# ALMDP Kaduna M&E Portfolio
# Results Framework Generator
# ============================================================

OUTPUT = Path("01_Project_Design/Results_Framework.xlsx")


# ============================================================
# Workbook setup
# ============================================================

wb = Workbook()

ws = wb.active
ws.title = "Results Framework"


# ============================================================
# Headers
# ============================================================

headers = [
    "Result Level",
    "Result Code",
    "Results Statement",
    "Indicator",
    "Indicator Definition / Calculation",
    "Baseline",
    "Baseline Year",
    "Target 2030",
    "Data Source",
    "Data Collection Method",
    "Frequency",
    "Responsible",
    "Disaggregation",
    "Assumptions / Risks",
    "Notes",
]

ws.append(headers)


# ============================================================
# Results Framework Indicators
# 2 Impact + 12 Outcome + 10 Output = 24
# ============================================================

rows = [

    # ========================================================
    # IMPACT INDICATORS
    # ========================================================

    [
        "Impact",
        "IMP-01",
        "Improved beneficiary economic well-being",
        "Average annual beneficiary income (₦/year)",
        "Total annual income across valid beneficiary observations divided by the number of valid observations.",
        420000,
        2026,
        600000,
        "Baseline/endline household survey",
        "Structured beneficiary survey",
        "Annual",
        "M&E Unit",
        "Sex, age group, LGA, beneficiary type, value chain",
        "Economic and market conditions remain sufficiently stable.",
        "Missing income values will not automatically be treated as zero.",
    ],

    [
        "Impact",
        "IMP-02",
        "Improved beneficiary economic resilience",
        "Beneficiaries reporting improved economic resilience (%)",
        "Number of beneficiaries meeting the defined improved-resilience criterion divided by valid assessed beneficiaries × 100.",
        "34%",
        2026,
        "65%",
        "Household survey",
        "Structured resilience assessment",
        "Annual",
        "M&E Unit",
        "Sex, age group, LGA, beneficiary type",
        "Beneficiaries remain exposed to manageable levels of economic and climate shocks.",
        "Higher percentage represents better performance.",
    ],


    # ========================================================
    # OUTCOME INDICATORS
    # ========================================================

    [
        "Outcome",
        "OUT-01",
        "Increased agricultural productivity",
        "Average crop productivity (t/ha)",
        "Quantity harvested in tonnes divided by cultivated area in hectares.",
        "1.80 t/ha",
        2026,
        "2.50 t/ha",
        "Farmer survey and production records",
        "Beneficiary survey / record review",
        "Annual",
        "Agriculture Technical Team / M&E",
        "LGA, value chain, sex, age group",
        "Inputs, weather and production conditions remain sufficiently favourable.",
        "Valid area and harvest observations required.",
    ],

    [
        "Outcome",
        "OUT-02",
        "Increased adoption of recommended practices",
        "Farmers adopting at least three recommended practices (%)",
        "Number of assessed farmers adopting at least three defined recommended practices divided by valid assessed farmers × 100.",
        "32%",
        2026,
        "65%",
        "Farmer survey",
        "Structured agricultural practices module",
        "Annual",
        "Agriculture Technical Team / M&E",
        "Sex, age group, LGA, value chain",
        "Farmers have access to inputs, services and technical information.",
        "Adoption threshold = at least 3 practices.",
    ],

    [
        "Outcome",
        "OUT-03",
        "Improved agricultural knowledge",
        "Trained farmers demonstrating adequate knowledge (%)",
        "Number of assessed trained farmers scoring at least 7 out of 10 knowledge questions divided by valid assessed trained farmers × 100.",
        "28%",
        2026,
        "75%",
        "Knowledge assessment",
        "Structured knowledge test",
        "Annual",
        "M&E Unit / Training Team",
        "Sex, age group, LGA, value chain",
        "Training content is relevant and participants engage with learning activities.",
        "Adequate knowledge = score ≥7/10.",
    ],

    [
        "Outcome",
        "OUT-04",
        "Improved market participation",
        "Farmers selling through structured markets (%)",
        "Number of valid farmer respondents selling through defined structured market channels divided by valid farmers assessed × 100.",
        "26%",
        2026,
        "55%",
        "Farmer survey",
        "Market access module",
        "Annual",
        "Market Access Team / M&E",
        "Sex, age group, LGA, value chain",
        "Buyers and structured market channels remain accessible.",
        "Structured market definition will be documented in the indicator reference sheet.",
    ],

    [
        "Outcome",
        "OUT-05",
        "Improved buyer access",
        "Beneficiaries reporting improved buyer access (%)",
        "Number of beneficiaries reporting improved access to buyers divided by valid assessed beneficiaries × 100.",
        "31%",
        2026,
        "70%",
        "Beneficiary survey",
        "Structured market access assessment",
        "Annual",
        "Market Access Team / M&E",
        "Sex, age group, LGA, beneficiary type, value chain",
        "Market linkage activities result in meaningful buyer connections.",
        "Self-reported outcome supported by market linkage records.",
    ],

    [
        "Outcome",
        "OUT-06",
        "Reduced post-harvest losses",
        "Average post-harvest loss rate (%)",
        "Quantity lost divided by quantity harvested × 100.",
        "22%",
        2026,
        "12%",
        "Farmer survey and production records",
        "Production and post-harvest assessment",
        "Annual",
        "Agriculture Technical Team / M&E",
        "LGA, value chain, sex, age group",
        "Improved storage, handling and transport practices are available.",
        "Lower value represents better performance.",
    ],

    [
        "Outcome",
        "OUT-07",
        "Improved agricultural MSME performance",
        "Supported MSMEs reporting increased annual revenue (%)",
        "Number of supported MSMEs reporting increased annual revenue compared with the defined comparison period divided by valid assessed supported MSMEs × 100.",
        "0%",
        2026,
        "60%",
        "MSME survey and business records",
        "Business survey / record review",
        "Annual",
        "MSME Team / M&E",
        "Sex, age group, LGA, business type",
        "MSMEs remain operational and market conditions support business activity.",
        "Baseline is 0% because project-supported MSMEs have not yet received project support.",
    ],

    [
        "Outcome",
        "OUT-08",
        "Improved MSME profitability",
        "Average MSME annual profit (₦/year)",
        "Total annual profit across valid MSME observations divided by the number of valid observations.",
        350000,
        2026,
        550000,
        "MSME survey and business records",
        "Business survey / record review",
        "Annual",
        "MSME Team / M&E",
        "LGA, sex, age group, business type",
        "MSMEs maintain operations and access viable markets.",
        "Profit = annual sales/revenue less annual operating costs.",
    ],

    [
        "Outcome",
        "OUT-09",
        "Improved MSME financial management",
        "Supported MSMEs maintaining basic financial records (%)",
        "Number of supported MSMEs meeting the defined financial-record standard divided by valid assessed supported MSMEs × 100.",
        "24%",
        2026,
        "70%",
        "MSME survey / record verification",
        "Record review and business survey",
        "Annual",
        "MSME Team / M&E",
        "Sex, age group, LGA, business type",
        "MSMEs are willing and able to adopt record-keeping practices.",
        "Record standard includes defined basic business records.",
    ],

    [
        "Outcome",
        "OUT-10",
        "Increased participation of women",
        "Direct beneficiaries who are women (%)",
        "Number of unique direct beneficiaries who are women divided by total unique direct beneficiaries × 100.",
        "35%",
        2026,
        "45%",
        "Project beneficiary database",
        "Administrative records",
        "Quarterly",
        "M&E Unit",
        "LGA, beneficiary type, value chain",
        "Women have equitable access to project activities.",
        "Unique beneficiaries must be counted once.",
    ],

    [
        "Outcome",
        "OUT-11",
        "Increased participation of youth",
        "Direct beneficiaries aged 18–35 (%)",
        "Number of unique direct beneficiaries aged 18–35 divided by total unique direct beneficiaries × 100.",
        "22%",
        2026,
        "35%",
        "Project beneficiary database",
        "Administrative records",
        "Quarterly",
        "M&E Unit",
        "Sex, LGA, beneficiary type, value chain",
        "Young people can access project opportunities.",
        "Youth definition = age 18–35.",
    ],

    [
        "Outcome",
        "OUT-12",
        "Improved economic outcomes for women",
        "Women beneficiaries reporting increased agricultural/business income (%)",
        "Number of assessed women beneficiaries reporting increased agricultural/business income divided by valid assessed women beneficiaries × 100.",
        "30%",
        2026,
        "60%",
        "Beneficiary survey",
        "Structured income/outcome assessment",
        "Annual",
        "M&E Unit / Gender Focal Point",
        "LGA, age group, beneficiary type, value chain",
        "Women can convert project support into sustained economic activity.",
        "Applies to women beneficiaries with valid income/outcome data.",
    ],


    # ========================================================
    # OUTPUT INDICATORS
    # ========================================================

    [
        "Output",
        "OP-01",
        "Farmers trained",
        "Number of farmers trained",
        "Number of unique farmers completing defined project training activities.",
        0,
        2026,
        8000,
        "Training attendance records",
        "Attendance register / digital training records",
        "Monthly",
        "Training Team / M&E",
        "Sex, age group, LGA, value chain",
        "Training activities are delivered as planned.",
        "Unique farmer counting required.",
    ],

    [
        "Output",
        "OP-02",
        "Farmers completing climate-smart agriculture training",
        "Number of farmers completing CSA training",
        "Number of unique farmers completing the defined climate-smart agriculture training package.",
        0,
        2026,
        5000,
        "CSA training records",
        "Attendance and completion records",
        "Monthly",
        "Agriculture Technical Team / M&E",
        "Sex, age group, LGA, value chain",
        "CSA training package is delivered as scheduled.",
        "Unique farmer counting required.",
    ],

    [
        "Output",
        "OP-03",
        "Demonstration plots established",
        "Number of demonstration plots established",
        "Number of verified demonstration plots established according to project specifications.",
        0,
        2026,
        80,
        "Project activity records",
        "Site verification / monitoring checklist",
        "Quarterly",
        "Agriculture Technical Team",
        "LGA, value chain",
        "Suitable sites and technical resources are available.",
        "Plots require verification before counting.",
    ],

    [
        "Output",
        "OP-04",
        "Market-linkage events conducted",
        "Number of market-linkage events",
        "Number of verified market-linkage events conducted according to project specifications.",
        0,
        2026,
        40,
        "Event records",
        "Event reports / attendance records",
        "Quarterly",
        "Market Access Team",
        "LGA, value chain, event type",
        "Buyers and value-chain actors participate.",
        "Each verified event counted once.",
    ],

    [
        "Output",
        "OP-05",
        "Farmers receiving market information",
        "Number of farmers receiving market information",
        "Number of unique farmers receiving defined project market information services.",
        0,
        2026,
        6000,
        "Market information records",
        "Digital/administrative records",
        "Monthly",
        "Market Access Team / M&E",
        "Sex, age group, LGA, value chain",
        "Market information is available and distributed through functioning channels.",
        "Unique beneficiary counting required.",
    ],

    [
        "Output",
        "OP-06",
        "MSMEs receiving business support",
        "Number of MSMEs receiving business support",
        "Number of unique agricultural MSMEs receiving defined project business development support.",
        0,
        2026,
        2000,
        "MSME support records",
        "Administrative records / support logs",
        "Monthly",
        "MSME Team / M&E",
        "Sex, age group, LGA, business type",
        "MSMEs participate in project support activities.",
        "Unique MSME counting required.",
    ],

    [
        "Output",
        "OP-07",
        "MSMEs completing financial management training",
        "Number of MSMEs completing financial management training",
        "Number of unique MSMEs completing the defined financial management training package.",
        0,
        2026,
        1500,
        "Training records",
        "Attendance and completion records",
        "Monthly",
        "MSME Team / M&E",
        "Sex, age group, LGA, business type",
        "Training is delivered and MSMEs complete the required sessions.",
        "Unique MSME counting required.",
    ],

    [
        "Output",
        "OP-08",
        "Timely M&E implementation",
        "Scheduled M&E activities completed on time (%)",
        "Number of scheduled M&E activities completed by their approved deadline divided by total scheduled M&E activities × 100.",
        "0%",
        2026,
        "≥90%",
        "M&E workplan",
        "M&E activity tracker",
        "Monthly",
        "M&E Manager",
        "Activity type, reporting period",
        "Field and operational resources are available on schedule.",
        "Higher percentage represents better performance.",
    ],

    [
        "Output",
        "OP-09",
        "Timely reporting",
        "Required reports submitted on time (%)",
        "Number of required reports submitted by the approved deadline divided by total required reports × 100.",
        "0%",
        2026,
        "100%",
        "Project reporting tracker",
        "Report submission tracker",
        "Monthly",
        "M&E Manager / Project Management",
        "Report type, reporting period",
        "Reporting inputs are received on time.",
        "Higher percentage represents better performance.",
    ],

    [
        "Output",
        "OP-10",
        "Data quality assurance",
        "Sampled records passing DQA (%)",
        "Number of sampled records meeting the defined DQA pass criteria divided by total sampled records × 100.",
        "0%",
        2026,
        "≥95%",
        "DQA records",
        "Data Quality Assessment",
        "Quarterly",
        "M&E Manager",
        "LGA, data source, indicator",
        "Data collection and reporting systems maintain adequate quality.",
        "DQA assesses accuracy, completeness, consistency, timeliness and integrity.",
    ],
]


# ============================================================
# Validate indicator list before creating workbook
# ============================================================

expected_codes = (
    ["IMP-01", "IMP-02"]
    + [f"OUT-{i:02d}" for i in range(1, 13)]
    + [f"OP-{i:02d}" for i in range(1, 11)]
)

actual_codes = [row[1] for row in rows]

missing_codes = [code for code in expected_codes if code not in actual_codes]
unexpected_codes = [code for code in actual_codes if code not in expected_codes]
duplicate_codes = sorted(
    {code for code in actual_codes if actual_codes.count(code) > 1}
)

if missing_codes:
    raise ValueError(f"Missing indicator codes: {missing_codes}")

if unexpected_codes:
    raise ValueError(f"Unexpected indicator codes: {unexpected_codes}")

if duplicate_codes:
    raise ValueError(f"Duplicate indicator codes: {duplicate_codes}")

if len(rows) != 24:
    raise ValueError(f"Expected 24 indicators but found {len(rows)}")


# ============================================================
# Add indicators to worksheet
# ============================================================

for row in rows:
    ws.append(row)


# ============================================================
# Formatting
# ============================================================

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)

thin = Side(style="thin", color="D9E1F2")

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = Border(bottom=thin)


for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


# Freeze header
ws.freeze_panes = "A2"

# Enable filters
ws.auto_filter.ref = ws.dimensions


# Column widths
widths = {
    "A": 14,
    "B": 14,
    "C": 34,
    "D": 42,
    "E": 58,
    "F": 16,
    "G": 14,
    "H": 16,
    "I": 32,
    "J": 34,
    "K": 16,
    "L": 30,
    "M": 38,
    "N": 48,
    "O": 45,
}

for column, width in widths.items():
    ws.column_dimensions[column].width = width

ws.row_dimensions[1].height = 42


# ============================================================
# Framework Metadata
# ============================================================

meta = wb.create_sheet("Framework Metadata")

metadata = [
    ["Field", "Value"],
    ["Project", "Agricultural Livelihoods and MSME Development Project (ALMDP)"],
    ["Status", "Synthetic/Simulated Portfolio Project"],
    ["Geographic Focus", "Kaduna State, Nigeria"],
    ["Implementation Period", "January 2027 – December 2030"],
    ["Target Beneficiaries", "10,000"],
    ["Farmers", "8,000"],
    ["Agricultural MSMEs", "2,000"],
    ["Number of Indicators", "24"],
    ["Impact Indicators", "2"],
    ["Outcome Indicators", "12"],
    ["Output Indicators", "10"],
    ["Baseline Year", "2026"],
    ["Target Year", "2030"],
]

for row in metadata:
    meta.append(row)

for cell in meta[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

for row in meta.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

meta.column_dimensions["A"].width = 28
meta.column_dimensions["B"].width = 65


# ============================================================
# Results Summary
# ============================================================

summary = wb.create_sheet("Results Summary")

summary_rows = [
    ["Result Level", "Number of Indicators"],
    ["Impact", 2],
    ["Outcome", 12],
    ["Output", 10],
    ["Total", 24],
]

for row in summary_rows:
    summary.append(row)

for cell in summary[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

summary.column_dimensions["A"].width = 24
summary.column_dimensions["B"].width = 24


# ============================================================
# Save workbook
# ============================================================

OUTPUT.parent.mkdir(parents=True, exist_ok=True)

wb.save(OUTPUT)

print(f"Created: {OUTPUT}")
print(f"Indicators: {len(rows)}")
print("Indicator validation: PASSED")
print("Expected codes: 24")