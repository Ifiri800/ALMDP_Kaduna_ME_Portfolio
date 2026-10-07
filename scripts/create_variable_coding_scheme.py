from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FILE = (
    PROJECT_ROOT
    / "04_Data_Metadata"
    / "Variable_Coding_Scheme.xlsx"
)


CODING_SCHEME = [
    {
        "variable": "consent",
        "label": "Respondent consent",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "Not permitted",
        "notes": "Consent must be recorded before questionnaire administration.",
    },
    {
        "variable": "state",
        "label": "State",
        "code_type": "Categorical",
        "codes": "KD = Kaduna",
        "missing_code": "Not permitted",
        "notes": "Synthetic project location.",
    },
    {
        "variable": "lga",
        "label": "Local Government Area",
        "code_type": "Categorical",
        "codes": (
            "KDN = Kaduna North; "
            "KDS = Kaduna South; "
            "CHK = Chikun; "
            "IGB = Igabi; "
            "KAJ = Kajuru; "
            "KAC = Kachia; "
            "ZAR = Zaria"
        ),
        "missing_code": "Not permitted",
        "notes": "Only approved project LGAs are valid.",
    },
    {
        "variable": "beneficiary_type",
        "label": "Beneficiary classification",
        "code_type": "Categorical",
        "codes": "1 = Farmer; 2 = MSME",
        "missing_code": "Not permitted",
        "notes": "Determines questionnaire skip patterns.",
    },
    {
        "variable": "sex",
        "label": "Sex",
        "code_type": "Categorical",
        "codes": "1 = Female; 2 = Male",
        "missing_code": "99 = Not stated",
        "notes": "Used for inclusion and disaggregation.",
    },
    {
        "variable": "education",
        "label": "Highest education level",
        "code_type": "Ordinal",
        "codes": (
            "1 = No formal education; "
            "2 = Primary; "
            "3 = Secondary; "
            "4 = Post-secondary"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Higher code indicates higher education level.",
    },
    {
        "variable": "disability_status",
        "label": "Disability status",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used for inclusion disaggregation.",
    },
    {
        "variable": "age_group",
        "label": "Age group",
        "code_type": "Derived categorical",
        "codes": (
            "1 = 18-35; "
            "2 = 36-59; "
            "3 = 60+"
        ),
        "missing_code": "Derived from valid age",
        "notes": "Do not collect independently; derive from age.",
    },
    {
        "variable": "youth_status",
        "label": "Youth status",
        "code_type": "Derived binary",
        "codes": "1 = Age 18-35; 0 = Otherwise",
        "missing_code": "Derived from valid age",
        "notes": "Used for OUT-11.",
    },
    {
        "variable": "primary_value_chain",
        "label": "Primary agricultural value chain",
        "code_type": "Categorical",
        "codes": (
            "1 = Maize; "
            "2 = Rice; "
            "3 = Soybean; "
            "4 = Ginger; "
            "5 = Tomato; "
            "6 = Vegetables; "
            "7 = Sorghum; "
            "8 = Groundnut"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Farmer respondents only.",
    },
    {
        "variable": "farmer_group_member",
        "label": "Farmer group membership",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Farmer respondents only.",
    },
    {
        "variable": "extension_access",
        "label": "Extension service access",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Farmer respondents only.",
    },
    {
        "variable": "improved_seed",
        "label": "Improved seed adoption",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "soil_testing",
        "label": "Soil testing adoption",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "fertilizer_management",
        "label": "Recommended fertilizer management",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "integrated_pest_management",
        "label": "Integrated pest management",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "conservation_agriculture",
        "label": "Conservation agriculture",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "water_management",
        "label": "Water management",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "improved_storage",
        "label": "Improved storage",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "climate_smart_practice",
        "label": "Climate-smart practice",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Included in practices_adopted.",
    },
    {
        "variable": "primary_market_channel",
        "label": "Primary market channel",
        "code_type": "Categorical",
        "codes": (
            "1 = Farm gate; "
            "2 = Local market; "
            "3 = Aggregator; "
            "4 = Cooperative; "
            "5 = Processor; "
            "6 = Wholesaler; "
            "7 = Other"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Farmer respondents only.",
    },
    {
        "variable": "regular_buyer",
        "label": "Regular buyer",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Farmer respondents only.",
    },
    {
        "variable": "structured_market",
        "label": "Structured market participation",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used for OUT-04.",
    },
    {
        "variable": "market_information_received",
        "label": "Market information received",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Supports market information monitoring.",
    },
    {
        "variable": "market_information_source",
        "label": "Market information source",
        "code_type": "Categorical",
        "codes": (
            "1 = Extension agent; "
            "2 = Farmer group; "
            "3 = Radio; "
            "4 = Phone/SMS; "
            "5 = Digital platform; "
            "6 = Buyer/aggregator; "
            "7 = Other"
        ),
        "missing_code": "99 = Not applicable/not stated",
        "notes": "Applicable where market information was received.",
    },
    {
        "variable": "buyer_access_improved",
        "label": "Improved buyer access",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Retrospective baseline self-report.",
    },
    {
        "variable": "business_type",
        "label": "MSME business type",
        "code_type": "Categorical",
        "codes": (
            "1 = Processing; "
            "2 = Trading; "
            "3 = Aggregation; "
            "4 = Input supply; "
            "5 = Storage; "
            "6 = Transport; "
            "7 = Other"
        ),
        "missing_code": "99 = Not stated",
        "notes": "MSME respondents only.",
    },
    {
        "variable": "sales_records",
        "label": "Sales records maintained",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used in financial_record_count.",
    },
    {
        "variable": "expense_records",
        "label": "Expense records maintained",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used in financial_record_count.",
    },
    {
        "variable": "inventory_records",
        "label": "Inventory records maintained",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used in financial_record_count.",
    },
    {
        "variable": "cashflow_records",
        "label": "Cash flow records maintained",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used in financial_record_count.",
    },
    {
        "variable": "profit_loss_records",
        "label": "Profit and loss records maintained",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used in financial_record_count.",
    },
    {
        "variable": "business_bank_account",
        "label": "Business bank account",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Descriptive financial-management variable; not included in core five-record count.",
    },
    {
        "variable": "women_beneficiary",
        "label": "Women beneficiary status",
        "code_type": "Derived binary",
        "codes": "1 = Female; 0 = Male",
        "missing_code": "Derived from sex",
        "notes": "Used for OUT-10 and OUT-12.",
    },
    {
        "variable": "income_increased",
        "label": "Income increased",
        "code_type": "Binary",
        "codes": "1 = Yes; 0 = No",
        "missing_code": "99 = Not stated",
        "notes": "Used for OUT-12 among women beneficiaries.",
    },
    {
        "variable": "resilience_crop_failure",
        "label": "Ability to cope with crop failure",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "resilience_price_shock",
        "label": "Ability to cope with price shock",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "resilience_input_price",
        "label": "Ability to cope with input price increases",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "resilience_health_shock",
        "label": "Ability to cope with health shock",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "resilience_climate_shock",
        "label": "Ability to cope with climate shock",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "resilience_income_loss",
        "label": "Ability to cope with income loss",
        "code_type": "Ordinal",
        "codes": (
            "1 = Can cope without assistance; "
            "2 = Can cope with some difficulty; "
            "3 = Requires assistance; "
            "4 = Unable to cope"
        ),
        "missing_code": "99 = Not stated",
        "notes": "Reverse-coded during resilience calculation.",
    },
    {
        "variable": "knowledge_q1",
        "label": "Knowledge question 1 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q2",
        "label": "Knowledge question 2 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q3",
        "label": "Knowledge question 3 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q4",
        "label": "Knowledge question 4 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q5",
        "label": "Knowledge question 5 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q6",
        "label": "Knowledge question 6 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q7",
        "label": "Knowledge question 7 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q8",
        "label": "Knowledge question 8 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q9",
        "label": "Knowledge question 9 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
    {
        "variable": "knowledge_q10",
        "label": "Knowledge question 10 score",
        "code_type": "Binary score",
        "codes": "1 = Correct; 0 = Incorrect",
        "missing_code": "Not automatically coded",
        "notes": "Derived from raw response and answer key.",
    },
]


def create_workbook():
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Coding Scheme"

    headers = [
        "Variable",
        "Label / Description",
        "Code Type",
        "Codes / Categories",
        "Missing Code",
        "Notes",
    ]

    worksheet.append(headers)

    for item in CODING_SCHEME:
        worksheet.append([
            item["variable"],
            item["label"],
            item["code_type"],
            item["codes"],
            item["missing_code"],
            item["notes"],
        ])

    # Missing-value standards
    missing = workbook.create_sheet("Missing Value Standards")
    missing.append(["Code", "Meaning", "Treatment"])
    missing_rows = [
        ("N/A", "Not applicable", "Retain as not applicable; exclude from affected denominator."),
        ("DK", "Don't know", "Do not convert to zero; review during DQA."),
        ("RF", "Refused", "Do not convert to zero; retain as missing/refused."),
        ("NA", "Not answered", "Review during data quality checks."),
        ("99", "Not stated / missing categorical response", "Review and document before analysis."),
        ("Blank", "Missing numeric or unentered value", "Do not automatically convert to zero."),
    ]

    for row in missing_rows:
        missing.append(row)

    # Coding principles
    principles = workbook.create_sheet("Coding Principles")
    principles.append(["Principle", "Requirement"])

    principle_rows = [
        (
            "Standardization",
            "The same code will represent the same response category throughout the dataset.",
        ),
        (
            "Traceability",
            "Codes will be traceable to questionnaire response options and the Data Dictionary.",
        ),
        (
            "Derived variables",
            "Derived variables will be calculated from validated source variables rather than entered independently.",
        ),
        (
            "Binary variables",
            "Binary variables will use 1 for Yes/positive and 0 for No/negative unless otherwise documented.",
        ),
        (
            "Missing values",
            "Missing values will not automatically be converted to zero.",
        ),
        (
            "Invalid values",
            "Values outside approved codes will be flagged during validation.",
        ),
        (
            "Raw data",
            "Raw collected values will be preserved before cleaning and transformation.",
        ),
    ]

    for row in principle_rows:
        principles.append(row)

    # Metadata
    metadata = workbook.create_sheet("Coding Metadata")
    metadata.append(["Field", "Value"])
    metadata_rows = [
        ("Project", "Agricultural Livelihoods and MSME Development Project (ALMDP)"),
        ("Location", "Kaduna State, Nigeria"),
        ("Dataset", "Synthetic baseline dataset"),
        ("Purpose", "Standardize coding for data collection, validation, cleaning, and analysis"),
        ("Status", "Synthetic / Simulated Portfolio Project"),
        ("Related file", "04_Data_Metadata/Data_Dictionary.xlsx"),
    ]

    for row in metadata_rows:
        metadata.append(row)

    # Formatting
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    widths = [28, 42, 22, 65, 35, 60]

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for sheet in [missing, principles, metadata]:
        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[column_letter].width = min(
                max(
                    len(str(cell.value or "")) for cell in column_cells
                ) + 2,
                80,
            )

    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(OUTPUT_FILE)

    print(f"Created: {OUTPUT_FILE}")
    print(f"Coded variables: {len(CODING_SCHEME)}")


if __name__ == "__main__":
    create_workbook()