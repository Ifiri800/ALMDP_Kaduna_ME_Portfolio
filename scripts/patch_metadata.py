from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


PROJECT_ROOT = Path(__file__).resolve().parents[1]

DATA_DICTIONARY = (
    PROJECT_ROOT
    / "04_Data_Metadata"
    / "Data_Dictionary.xlsx"
)

CODING_SCHEME = (
    PROJECT_ROOT
    / "04_Data_Metadata"
    / "Variable_Coding_Scheme.xlsx"
)


def patch_data_dictionary():
    workbook = load_workbook(DATA_DICTIONARY)
    sheet = workbook["Data Dictionary"]

    existing = {
        row[0].value
        for row in sheet.iter_rows(min_row=2)
        if row[0].value
    }

    additions = [
        [
            "productivity_t_per_ha",
            "Crop productivity",
            "Numeric",
            "Continuous",
            "tonnes per hectare",
            "Non-negative numeric value",
            "Missing if area cultivated or harvested quantity is invalid",
            "Section 12 - Practice Adoption and Productivity Derivation",
            "OUT-01",
            "Yes",
            "quantity_harvested_t / area_cultivated_ha",
        ],
    ]

    for question_number in range(1, 11):
        additions.append(
            [
                f"knowledge_q{question_number}_response",
                f"Knowledge question {question_number} raw response",
                "Integer",
                "Nominal",
                "Response option",
                "1 = Option 1; 2 = Option 2; 3 = Option 3; 4 = Option 4",
                "Missing if respondent does not provide a valid response",
                "Section 21 - Agricultural Knowledge Assessment",
                "OUT-03",
                "No",
                "Raw response retained; compared with answer key to derive knowledge_qN",
            ]
        )

    added = 0

    for row in additions:
        variable = row[0]

        if variable not in existing:
            sheet.append(row)
            added += 1

    # Apply wrapping to newly added rows and all dictionary cells.
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(DATA_DICTIONARY)

    print(f"Data Dictionary additions: {added}")
    print(f"Data Dictionary variables: {sheet.max_row - 1}")


def patch_coding_scheme():
    workbook = load_workbook(CODING_SCHEME)
    sheet = workbook["Coding Scheme"]

    existing = {
        row[0].value
        for row in sheet.iter_rows(min_row=2)
        if row[0].value
    }

    answer_options = (
        "1 = Option 1; "
        "2 = Option 2; "
        "3 = Option 3; "
        "4 = Option 4"
    )

    added = 0

    # Raw knowledge response variables.
    for question_number in range(1, 11):
        variable = f"knowledge_q{question_number}_response"

        if variable not in existing:
            sheet.append(
                [
                    variable,
                    f"Knowledge question {question_number} raw response",
                    "Categorical response",
                    answer_options,
                    "Blank = Missing response",
                    (
                        "Raw response retained for traceability. "
                        "The response will be compared with the answer key "
                        "to derive the corresponding binary knowledge score."
                    ),
                ]
            )
            added += 1

    # Productivity is numeric and therefore does not require a categorical
    # code list, but documenting its derivation in the coding scheme improves
    # metadata traceability.
    if "productivity_t_per_ha" not in existing:
        sheet.append(
            [
                "productivity_t_per_ha",
                "Crop productivity",
                "Derived numeric",
                "Continuous non-negative value",
                "Not applicable",
                "Derived as quantity_harvested_t / area_cultivated_ha.",
            ]
        )
        added += 1

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(CODING_SCHEME)

    print(f"Coding Scheme additions: {added}")
    print(f"Coding Scheme variables: {sheet.max_row - 1}")


if __name__ == "__main__":
    patch_data_dictionary()
    patch_coding_scheme()